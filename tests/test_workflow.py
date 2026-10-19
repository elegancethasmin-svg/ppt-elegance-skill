import base64
import copy
import importlib.util
import io
import json
import os
from pathlib import Path
import struct
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch, MagicMock
import urllib.error
import zipfile
import zlib

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from common import read_json, validate, write_json
from build_prompts import compile_prompt
from generate_image import prepare, request_image, decode_image, NoRedirect

def png_pixel():
    def chunk(kind, body):
        return struct.pack('!I', len(body)) + kind + body + struct.pack('!I', zlib.crc32(kind+body) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('!2I5B',1,1,8,2,0,0,0)) + chunk(b'IDAT', zlib.compress(b'\0\xff\xff\xff')) + chunk(b'IEND', b'')

def config():
    return {'adapter':'images-json','base_url':'https://images.test.invalid/v1',
            'model':'test-model','api_key_env':'TEST_IMAGE_KEY','parameters':{'size':'1920x1080'}}

class SpecsAndPrompts(unittest.TestCase):
    def setUp(self):
        self.deck = read_json(ROOT/'examples/activity.json')

    def test_examples_and_all_styles(self):
        for example in (ROOT/'examples').glob('*.json'):
            validate(read_json(example))
        for style in (ROOT/'styles').glob('*.json'):
            d = copy.deepcopy(self.deck); d['style'] = style.stem
            validate(d)
            prompt = compile_prompt(d, d['slides'][0])
            self.assertIn(read_json(style)['name'], prompt)

    def test_invalid_page_contracts(self):
        for mutate in [lambda d: d.update(slide_count=99),
                       lambda d: d['slides'][0].update(page=True),
                       lambda d: d['slides'][0].update(title='not in copy'),
                       lambda d: d['slides'][0].update(visible_text='not a list'),
                       lambda d: d.update(style='../private'),
                       lambda d: d['slides'][0].update(preserve_pixels=True),
                       lambda d: d['slides'][0].update(image_box=[.9,.2,.5,.5]),
                       lambda d: d['slides'][0].update(reviewed='yes')]:
            d = copy.deepcopy(self.deck); mutate(d)
            with self.assertRaises(ValueError): validate(d)

    def test_only_selected_style_and_visible_copy(self):
        p = self.deck['slides'][0]
        p.update(speaker_notes='PRIVATE_NOTES', source_basis=['PRIVATE_FILE'],
                 key_message='HIDDEN_KEY_MESSAGE', source_image='PRIVATE_IMAGE')
        prompt = compile_prompt(self.deck,p)
        for text in p['visible_text']: self.assertIn(text,prompt)
        for text in ['PRIVATE_NOTES','PRIVATE_FILE','HIDDEN_KEY_MESSAGE','PRIVATE_IMAGE','石墨黑']:
            self.assertNotIn(text,prompt)

    def test_compile_cli_and_no_overwrite(self):
        with tempfile.TemporaryDirectory() as temp:
            cmd=[sys.executable,str(ROOT/'scripts/build_prompts.py'),str(ROOT/'examples/activity.json'),'--out-dir',temp+'/prompts']
            one=subprocess.run(cmd,capture_output=True)
            self.assertEqual(one.returncode,0,one.stderr)
            self.assertEqual(len(list(Path(temp+'/prompts').glob('P*.txt'))),3)
            two=subprocess.run(cmd,capture_output=True)
            self.assertNotEqual(two.returncode,0)

class AdapterTests(unittest.TestCase):
    def test_request_contract_and_decode(self):
        url, payload, timeout, env=prepare(config(),'exact prompt')
        self.assertEqual(url,'https://images.test.invalid/v1/images/generations')
        self.assertEqual(payload['prompt'],'exact prompt')
        self.assertEqual(payload['n'],1)
        response=MagicMock()
        response.__enter__.return_value.read.return_value=json.dumps({'data':[{'b64_json':base64.b64encode(png_pixel()).decode()}]}).encode()
        opener=MagicMock(); opener.open.return_value=response
        with patch('urllib.request.build_opener',return_value=opener):
            self.assertEqual(request_image(url,payload,timeout,'test-secret'),png_pixel())
        req=opener.open.call_args.args[0]
        self.assertEqual(req.get_header('Authorization'),'Bearer test-secret')
        self.assertEqual(json.loads(req.data),payload)

    def test_config_rejections(self):
        for mutation in [{'base_url':'http://images.test.invalid'},
                         {'base_url':'https://user:password@images.test.invalid'},
                         {'base_url':'https://images.test.invalid?token=example'},
                         {'parameters':{'model':'override'}}, {'timeout_seconds':0},
                         {'model':'your-image-model'}, {'api_key_env':'bad name'},
                         {'parameters':{'output_format':'jpeg'}}]:
            c=config();c.update(mutation)
            with self.assertRaises(ValueError):prepare(c,'prompt')

    def test_no_key_or_network_on_dry_run(self):
        from generate_image import main
        with tempfile.TemporaryDirectory() as temp:
            c=Path(temp)/'config.json'; write_json(c,config())
            p=Path(temp)/'prompt.txt';p.write_text('hello',encoding='utf-8')
            args=['script','--config',str(c),'--prompt-file',str(p),'--out',temp+'/page.png','--dry-run']
            env_get=os.environ.get
            def guarded_get(name, default=None):
                if name=='TEST_IMAGE_KEY':raise AssertionError('key was read')
                return env_get(name, default)
            with patch.object(sys,'argv',args),patch('os.environ.get',side_effect=guarded_get),patch('urllib.request.build_opener',side_effect=AssertionError('network used')),patch('sys.stdout',new=io.StringIO()):
                main()

    def test_errors_hide_provider_body_and_no_retry(self):
        opener=MagicMock()
        opener.open.side_effect=urllib.error.HTTPError('https://images.test.invalid',401,'provider-secret',{},io.BytesIO(b'private detail'))
        with patch('urllib.request.build_opener',return_value=opener):
            with self.assertRaises(RuntimeError) as caught:request_image('https://images.test.invalid',{},1,'test-secret')
        self.assertNotIn('provider-secret',str(caught.exception))
        self.assertNotIn('private detail',str(caught.exception))
        self.assertEqual(opener.open.call_count,1)
        self.assertIsNone(NoRedirect().redirect_request(None,None,302,'',{},'https://elsewhere.invalid'))

    def test_malformed_responses(self):
        for obj in [{'data':[{'url':'https://example.com/image.png'}]}, {'task_id':'id'},
                    {'data':[{'b64_json':'!!!'}]}, {'data':[{'b64_json':base64.b64encode(b'not png').decode()}]}]:
            with self.assertRaises(ValueError):decode_image(json.dumps(obj).encode())

HAS_PACKAGING=all(importlib.util.find_spec(m) is not None for m in ['PIL','pptx','reportlab'])

@unittest.skipUnless(HAS_PACKAGING,'Install requirements.txt to test packaging')
class PackagingTests(unittest.TestCase):
    def test_image_notes_pdf_zip_and_original(self):
        from PIL import Image
        from pptx import Presentation
        from package_deck import package
        with tempfile.TemporaryDirectory() as temp:
            base=Path(temp)
            Image.new('RGB',(1600,900),'navy').save(base/'page.png')
            Image.new('RGB',(40,80),'red').save(base/'evidence.png')
            deck=read_json(ROOT/'examples/technical.json')
            deck['slides']=deck['slides'][:1];deck['slide_count']=1
            p=deck['slides'][0];p.update(image='page.png',reviewed=True,preserve_pixels=True,
                source_image='evidence.png',image_box=[.1,.2,.4,.5],speaker_notes='Known speaker note')
            write_json(base/'spec.json',deck)
            result=package(base/'spec.json',base/'out')
            prs=Presentation(result/'presentation.pptx')
            self.assertEqual(len(prs.slides),1)
            self.assertIn('Known speaker note',prs.slides[0].notes_slide.notes_text_frame.text)
            self.assertIn(b'/Count 1',(result/'presentation.pdf').read_bytes())
            with zipfile.ZipFile(result/'pages.zip') as z:self.assertEqual(z.namelist(),['P01.png'])
            im=Image.open(result/'pages/P01.png')
            self.assertEqual(im.size,(1920,1080))
            self.assertEqual(im.getpixel((576,486)),(255,0,0))
            self.assertEqual(im.getpixel((195,220)),(255,255,255))
            self.assertTrue(read_json(result/'manifest.json')['post_composite_review_required'])

    def test_refuse_unreviewed_and_non_widescreen(self):
        from PIL import Image
        from package_deck import package
        with tempfile.TemporaryDirectory() as temp:
            base=Path(temp);Image.new('RGB',(100,100)).save(base/'page.png')
            deck=read_json(ROOT/'examples/technical.json');deck['slides']=deck['slides'][:1];deck['slide_count']=1
            deck['slides'][0].update(image='page.png',reviewed=False)
            write_json(base/'spec.json',deck)
            with self.assertRaises(ValueError):package(base/'spec.json',base/'out')
            deck['slides'][0]['reviewed']=True;write_json(base/'spec.json',deck)
            with self.assertRaises(ValueError):package(base/'spec.json',base/'out')
            self.assertFalse((base/'out').exists())

if __name__=='__main__':unittest.main()
