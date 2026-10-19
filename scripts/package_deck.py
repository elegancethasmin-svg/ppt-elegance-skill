"""Package reviewed image slides; optional original-image insertion, no network."""
import argparse
import hashlib
import zipfile
from pathlib import Path
from common import read_json, validate, write_json

def package(spec_path, out_dir):
    from PIL import Image, ImageOps
    from pptx import Presentation
    from pptx.util import Inches
    from reportlab.pdfgen import canvas
    from reportlab.lib.utils import ImageReader

    spec_path = Path(spec_path).resolve()
    deck = validate(read_json(spec_path))
    out = Path(out_dir)
    if out.exists() and any(out.iterdir()):
        raise ValueError('Use a new or empty output directory')
    prepared = []
    # Check every input before writing any output; all paths are relative to the spec.
    for p in deck['slides']:
        if p.get('reviewed') is not True:
            raise ValueError(f"Page {p['page']}: inspect the image, then set reviewed to true")
        if not p.get('image'):
            raise ValueError(f"Page {p['page']}: image path is required")
        src = spec_path.parent / p['image']
        with Image.open(src) as original:
            im = ImageOps.exif_transpose(original).convert('RGB')
        if abs(im.width / im.height - 16/9) > 0.005:
            raise ValueError(f"Page {p['page']}: image is not 16:9; review and recompose or letterbox it explicitly")
        im = im.resize((1920, 1080), Image.Resampling.LANCZOS)
        if p.get('preserve_pixels'):
            with Image.open(spec_path.parent / p['source_image']) as original:
                evidence = ImageOps.exif_transpose(original).convert('RGBA')
            x, y, w, h = p['image_box']
            area = (round(w*1920), round(h*1080))
            evidence.thumbnail(area, Image.Resampling.LANCZOS)
            # Letterbox within the designated slot; no crop, stretch, or synthesized content.
            slot = Image.new('RGB', area, p.get('image_box_color', '#FFFFFF'))
            slot.paste(evidence, ((area[0]-evidence.width)//2, (area[1]-evidence.height)//2), evidence)
            im.paste(slot, (round(x*1920), round(y*1080)))
        prepared.append(im)
    out.mkdir(parents=True, exist_ok=True)
    image_dir = out / 'pages'
    image_dir.mkdir(exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(40/3), Inches(7.5)
    prs.core_properties.title = deck['title']
    prs.core_properties.author = ''
    pdf = canvas.Canvas(str(out/'presentation.pdf'), pagesize=(960, 540))
    pdf.setTitle(deck['title'])
    sheet = Image.new('RGB', (1280, ((len(prepared)+3)//4)*188), '#E9EDF3')
    manifest = []
    for n, (p, im) in enumerate(zip(deck['slides'], prepared), 1):
        path = image_dir/f'P{n:02d}.png'
        im.save(path)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(str(path), 0, 0, width=prs.slide_width, height=prs.slide_height)
        slide.notes_slide.notes_text_frame.text = '\n\n'.join(filter(None, [
            p['title'], p.get('speaker_notes', ''), 'Sources: '+'; '.join(p.get('source_basis', [])),
            'Evidence status: '+p['evidence_status']]))
        pdf.drawImage(ImageReader(str(path)), 0, 0, width=960, height=540)
        pdf.showPage()
        sheet.paste(im.resize((312, 176), Image.Resampling.LANCZOS), (((n-1)%4)*320+4, ((n-1)//4)*188+4))
        manifest.append({'page':n, 'file':f'pages/{path.name}', 'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),
                         'original_inserted': bool(p.get('preserve_pixels'))})
    pdf.save()
    prs.save(out/'presentation.pptx')
    sheet.save(out/'overview.jpg', quality=92)
    with zipfile.ZipFile(out/'pages.zip', 'w', zipfile.ZIP_DEFLATED) as z:
        for p in sorted(image_dir.glob('P*.png')):
            z.write(p, p.name)
    write_json(out/'manifest.json', {'slide_count':len(prepared), 'slides':manifest,
        'output_mode':'image-based', 'post_composite_review_required':any(p.get('preserve_pixels') for p in deck['slides'])})
    return out

def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('spec')
    p.add_argument('--out-dir', required=True)
    args = p.parse_args()
    try:
        package(args.spec, args.out_dir)
    except ImportError:
        p.exit(1, 'Install optional packaging dependencies: python -m pip install -r requirements.txt\n')
    except (ValueError, OSError) as exc:
        p.exit(1, f'Packaging failed: {exc}\n')
    print('Packaged image-based PPTX, PDF, PNG ZIP, overview and manifest. Inspect final pages, especially inserted originals.')

if __name__ == '__main__':
    main()
